"""Build editable Office files from the P008 Markdown sources.

Uses python-docx and python-pptx when they are already installed.
Does not claim a PDF was produced.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches as PInches
from pptx.util import Pt as PPt

ROOT = Path(__file__).resolve().parent


def add_report(source: Path, target: Path) -> None:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            index += 1
            block: list[str] = []
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            paragraph = document.add_paragraph("\n".join(block))
            for run in paragraph.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            index += 1
            continue
        if line.startswith("| ") and index + 1 < len(lines) and lines[index + 1].startswith("| ---"):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                if not lines[index].startswith("| ---"):
                    cells = [cell.strip() for cell in lines[index].strip("|").split("|")]
                    rows.append(cells)
                index += 1
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for row_index, cells in enumerate(rows):
                for col_index, cell in enumerate(cells):
                    table.rows[row_index].cells[col_index].text = cell
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.strip():
            document.add_paragraph(line.strip())
        index += 1
    document.save(target)


def add_slides(source: Path, target: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = PInches(13.333)
    presentation.slide_height = PInches(7.5)
    chunks = re.split(r"(?m)^## ", source.read_text(encoding="utf-8"))
    for chunk in chunks:
        title, _, body = chunk.partition("\n")
        title = title.strip()
        if not title or title.startswith("#"):
            continue
        layout = presentation.slide_layouts[1]
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.clear()
        first = True
        for line in body.splitlines():
            text = line.strip()
            if not text or text.startswith("#"):
                continue
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            paragraph.text = text[2:].strip() if text.startswith("- ") else text
            paragraph.level = 0
            paragraph.font.size = PPt(20)
    presentation.save(target)


if __name__ == "__main__":
    add_report(ROOT / "report" / "BOWLBYTE_PROJECT_REPORT.md", ROOT / "report" / "BOWLBYTE_PROJECT_REPORT.docx")
    add_slides(ROOT / "presentation" / "BOWLBYTE_SLIDES.md", ROOT / "presentation" / "BOWLBYTE_SLIDES.pptx")
    print("wrote docx and pptx")
